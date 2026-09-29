from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Create a presentation object
prs = Presentation()

# Define colors for the technology theme (blues)
primary_color = RGBColor(25, 118, 210)
accent_color = RGBColor(100, 181, 246)
background_color = RGBColor(227, 242, 253)
text_color = RGBColor(33, 33, 33)
white = RGBColor(255, 255, 255)

# Slide 1: Title Slide
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
fill = slide.background.fill
fill.solid()
fill.fore_color.rgb = background_color

title = slide.shapes.title
title.text = "Artificial Intelligence"
title.text_frame.paragraphs[0].font.color.rgb = primary_color
title.text_frame.paragraphs[0].font.size = Pt(44)
title.text_frame.paragraphs[0].font.bold = True

subtitle = slide.placeholders[1]
subtitle.text = "Understanding the Future"
subtitle.text_frame.paragraphs[0].font.color.rgb = text_color
subtitle.text_frame.paragraphs[0].font.size = Pt(24)

# Slide 2: What is AI?
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
fill = slide.background.fill
fill.solid()
fill.fore_color.rgb = background_color

title = slide.shapes.title
title.text = "What is Artificial Intelligence?"
title.text_frame.paragraphs[0].font.color.rgb = primary_color
title.text_frame.paragraphs[0].font.size = Pt(36)
title.text_frame.paragraphs[0].font.bold = True

body_shape = slide.shapes.placeholders[1]
tf = body_shape.text_frame
tf.clear()

p = tf.add_paragraph()
p.text = "AI refers to the simulation of human intelligence in machines that are programmed to think like humans and mimic their actions."
p.font.size = Pt(20)
p.font.color.rgb = text_color
p.level = 0

p = tf.add_paragraph()
p.text = "It involves creating algorithms and systems that can perform tasks typically requiring human intelligence, such as learning, problem-solving, decision-making, and perception."
p.font.size = Pt(20)
p.font.color.rgb = text_color
p.level = 1

# Slide 3: Types of AI
slide = prs.slides.add_slide(prs.slide_layouts[1])
fill = slide.background.fill
fill.solid()
fill.fore_color.rgb = background_color

title = slide.shapes.title
title.text = "Types of AI"
title.text_frame.paragraphs[0].font.color.rgb = primary_color
title.text_frame.paragraphs[0].font.size = Pt(36)
title.text_frame.paragraphs[0].font.bold = True

body_shape = slide.shapes.placeholders[1]
tf = body_shape.text_frame
tf.clear()

p = tf.add_paragraph()
p.text = "Narrow AI (Weak AI): Designed and trained for a specific task."
p.font.size = Pt(20)
p.font.color.rgb = text_color
p.level = 0

p = tf.add_paragraph()
p.text = "General AI (Strong AI): Hypothetical AI with human-like cognitive abilities."
p.font.size = Pt(20)
p.font.color.rgb = text_color
p.level = 0

p = tf.add_paragraph()
p.text = "Superintelligence AI: AI that surpasses human intelligence and ability."
p.font.size = Pt(20)
p.font.color.rgb = text_color
p.level = 0

# Slide 4: Applications of AI
slide = prs.slides.add_slide(prs.slide_layouts[1])
fill = slide.background.fill
fill.solid()
fill.fore_color.rgb = background_color

title = slide.shapes.title
title.text = "Applications of AI"
title.text_frame.paragraphs[0].font.color.rgb = primary_color
title.text_frame.paragraphs[0].font.size = Pt(36)
title.text_frame.paragraphs[0].font.bold = True

body_shape = slide.shapes.placeholders[1]
tf = body_shape.text_frame
tf.clear()

data = [
    ["Application Area", "Examples"],
    ["Healthcare", "Diagnosis, Drug Discovery, Personalized Medicine"],
    ["Finance", "Fraud Detection, Algorithmic Trading, Credit Scoring"],
    ["Transportation", "Self-Driving Cars, Traffic Management, Route Optimization"],
    ["Entertainment", "Recommendation Systems, Game AI, Content Generation"],
    ["Customer Service", "Chatbots, Virtual Assistants, Sentiment Analysis"]
]
rows, cols = len(data), len(data[0])
left, top, width, height = Inches(0.5), Inches(1.5), Inches(9), Inches(0.4 * rows)
table = slide.shapes.add_table(rows, cols, left, top, width, height).table

# Set table header style
for j in range(cols):
    cell = table.cell(0, j)
    cell.text = data[0][j]
    cell.fill.solid()
    cell.fill.fore_color.rgb = primary_color
    run = cell.text_frame.paragraphs[0].runs[0]
    run.font.color.rgb = white
    run.font.bold = True
    cell.text_frame.paragraphs[0].font.size = Pt(14)

# Fill table data and apply alternating row colors
for i, row_data in enumerate(data[1:], start=1):
    for j, val in enumerate(row_data):
        cell = table.cell(i, j)
        cell.text = str(val)
        cell.text_frame.paragraphs[0].font.size = Pt(12)
        cell.text_frame.paragraphs[0].font.color.rgb = text_color
        if i % 2 == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = white
        else:
            cell.fill.solid()
            cell.fill.fore_color.rgb = background_color

# Slide 5: The Future of AI
slide = prs.slides.add_slide(prs.slide_layouts[1])
fill = slide.background.fill
fill.solid()
fill.fore_color.rgb = background_color

title = slide.shapes.title
title.text = "The Future of AI"
title.text_frame.paragraphs[0].font.color.rgb = primary_color
title.text_frame.paragraphs[0].font.size = Pt(36)
title.text_frame.paragraphs[0].font.bold = True

body_shape = slide.shapes.placeholders[1]
tf = body_shape.text_frame
tf.clear()

p = tf.add_paragraph()
p.text = "AI is rapidly evolving and is expected to play an even more significant role in our lives."
p.font.size = Pt(20)
p.font.color.rgb = text_color
p.level = 0

p = tf.add_paragraph()
p.text = "Potential advancements include more sophisticated machine learning, improved natural language processing, and enhanced human-AI collaboration."
p.font.size = Pt(20)
p.font.color.rgb = text_color
p.level = 1

# Add a decorative shape
left, top, width, height = Inches(8.5), Inches(4.5), Inches(1.5), Inches(1.5)
shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, width, height)
fill = shape.fill
fill.solid()
fill.fore_color.rgb = accent_color
line = shape.line
line.color.rgb = accent_color
line.width = Pt(2)

# Save the presentation
prs.save("artificial_intelligence_presentation.pptx")